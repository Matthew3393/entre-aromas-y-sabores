#!/usr/bin/env python3
"""
viandas_engine.py - Motor de gestión de viandas, precios y pedidos con exportación a Excel.
"""

import os
import re
import json
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path

# Rutas predeterminadas
BASE_DIR = Path(__file__).resolve().parent
CATALOGO_FILE = BASE_DIR / "catalogo.json"
PEDIDOS_FILE = BASE_DIR / "pedidos.json"
EXCEL_OUTPUT = BASE_DIR / "viandas.xlsx"
USUARIOS_FILE = BASE_DIR / "usuarios.json"
DESKTOP_DIR = Path("/home/matthew/Escritorio")
DESKTOP_EXCEL = DESKTOP_DIR / "viandas.xlsx"

DIAS_ORDEN = {
    "lunes": 1,
    "martes": 2,
    "miercoles": 3,
    "miércoles": 3,
    "jueves": 4,
    "viernes": 5,
    "sabado": 6,
    "sábado": 6,
    "domingo": 7
}

def normalize_text(text: str) -> str:
    """Elimina tildes, signos de puntuación extra y pasa a minúsculas."""
    if not text:
        return ""
    # Descomponer caracteres con acentos
    text = unicodedata.normalize('NFD', text)
    text = "".join(c for c in text if unicodedata.category(c) != 'Mn')
    text = re.sub(r'[^a-zA-Z0-9\s]', ' ', text)
    return " ".join(text.lower().split())

class GestorUsuarios:
    def __init__(self, filepath=USUARIOS_FILE):
        self.filepath = Path(filepath)
        self.usuarios = []
        self.load()

    def load(self):
        if self.filepath.exists():
            try:
                with open(self.filepath, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self.usuarios = data.get("usuarios", [])
            except Exception as e:
                print(f"[!] Error al cargar usuarios: {e}")
                self.usuarios = []
        else:
            self.usuarios = []

    def save(self):
        with open(self.filepath, "w", encoding="utf-8") as f:
            json.dump({"usuarios": self.usuarios}, f, indent=2, ensure_ascii=False)

    def get_by_id(self, user_id):
        try:
            uid = int(user_id)
        except (ValueError, TypeError):
            return None
        for u in self.usuarios:
            if u.get("id") == uid:
                return u
        return None

    def get_by_whatsapp(self, whatsapp: str):
        wa_digits = re.sub(r'\D', '', str(whatsapp))
        if not wa_digits:
            return None
        for u in self.usuarios:
            u_digits = re.sub(r'\D', '', str(u.get("whatsapp", "")))
            if u_digits and u_digits == wa_digits:
                return u
        return None

    def registrar_o_actualizar(self, nombre: str, whatsapp: str, direccion: str = "", notas: str = ""):
        wa_digits = re.sub(r'\D', '', str(whatsapp))
        nombre_clean = nombre.strip().title()

        usuario = None
        if wa_digits:
            usuario = self.get_by_whatsapp(wa_digits)

        # Si no lo encuentra por WhatsApp, probar por nombre exacto
        if not usuario and nombre_clean:
            for u in self.usuarios:
                if normalize_text(u.get("nombre", "")) == normalize_text(nombre_clean):
                    usuario = u
                    break

        ahora = datetime.now().strftime("%Y-%m-%d %H:%M")

        if usuario:
            # Actualizar datos
            if nombre_clean:
                usuario["nombre"] = nombre_clean
            if whatsapp:
                usuario["whatsapp"] = whatsapp.strip()
            if direccion:
                usuario["direccion"] = direccion.strip()
            if notas:
                usuario["notas_frecuentes"] = notas.strip()
            usuario["ultimo_acceso"] = ahora
            usuario["pedidos_count"] = usuario.get("pedidos_count", 0) + 1
        else:
            next_id = max([u.get("id", 0) for u in self.usuarios], default=0) + 1
            usuario = {
                "id": next_id,
                "nombre": nombre_clean,
                "whatsapp": whatsapp.strip(),
                "direccion": direccion.strip(),
                "notas_frecuentes": notas.strip(),
                "pedidos_count": 1,
                "fecha_registro": ahora,
                "ultimo_acceso": ahora
            }
            self.usuarios.append(usuario)

        self.save()
        return usuario

class Catalogo:
    def __init__(self, filepath=CATALOGO_FILE):
        self.filepath = Path(filepath)
        self.items = []
        self.load()

    def load(self):
        if self.filepath.exists():
            try:
                with open(self.filepath, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self.items = data.get("items", [])
            except Exception as e:
                print(f"[!] Error al cargar catálogo: {e}")
                self.items = []
        else:
            self.items = []

    def save(self):
        with open(self.filepath, "w", encoding="utf-8") as f:
            json.dump({"items": self.items}, f, indent=2, ensure_ascii=False)

    def add_or_update_item(self, nombre: str, precio: float, aliases: list = None, foto: str = "", item_id: int = None, dias: list = None, sabores: list = None, descripcion: str = "", tipo: str = ""):
        aliases = aliases or []
        dias = dias or ["Todos"]
        sabores = sabores or []
        norm_name = normalize_text(nombre)
        found = False
        if item_id:
            try:
                iid = int(item_id)
                for it in self.items:
                    if it.get("id") == iid:
                        it["nombre"] = nombre.strip()
                        it["precio"] = float(precio)
                        if aliases:
                            it["aliases"] = aliases
                        if foto:
                            it["foto"] = foto
                        it["dias"] = dias
                        it["sabores"] = sabores
                        if descripcion:
                            it["descripcion"] = descripcion
                        if tipo:
                            it["tipo"] = tipo
                        found = True
                        break
            except (ValueError, TypeError):
                pass

        if not found:
            for it in self.items:
                if normalize_text(it["nombre"]) == norm_name:
                    it["precio"] = float(precio)
                    if foto:
                        it["foto"] = foto
                    for a in aliases:
                        if a not in it.get("aliases", []):
                            it.setdefault("aliases", []).append(a)
                    it["dias"] = dias
                    it["sabores"] = sabores
                    if descripcion:
                        it["descripcion"] = descripcion
                    if tipo:
                        it["tipo"] = tipo
                    found = True
                    break

        if not found:
            next_id = max([it.get("id", 0) for it in self.items], default=0) + 1
            if not aliases:
                aliases = [normalize_text(nombre)]
            self.items.append({
                "id": next_id,
                "nombre": nombre.strip(),
                "precio": float(precio),
                "aliases": aliases,
                "foto": foto or "",
                "dias": dias,
                "sabores": sabores,
                "descripcion": descripcion,
                "tipo": tipo
            })
        self.save()

    def delete_item(self, item_id: int):
        try:
            iid = int(item_id)
            initial_count = len(self.items)
            self.items = [it for it in self.items if it.get("id") != iid]
            if len(self.items) < initial_count:
                self.save()
                return True
            return False
        except (ValueError, TypeError):
            return False

    def set_items(self, new_items: list):
        """Reemplaza la lista completa de platos y precios."""
        self.items = []
        for i, item in enumerate(new_items, 1):
            aliases = item.get("aliases", [])
            if not aliases:
                aliases = [normalize_text(item["nombre"])]
            self.items.append({
                "id": i,
                "nombre": item["nombre"],
                "precio": float(item["precio"]),
                "aliases": aliases
            })
        self.save()

    def parse_and_set_menu_text(self, text: str):
        """
        Interpreta un bloque de texto con platos y precios, por ejemplo:
        Pollo con papas: 1500
        Milanesa con puré - 1800 (alias: mila, milanga)
        Tarta de jamon y queso 1200
        $1100 Pastas
        """
        lines = [l.strip() for l in text.strip().splitlines() if l.strip()]
        new_items = []
        for line in lines:
            # Buscar precio (número entero o decimal, con o sin $)
            m_price = re.search(r'\$?\s*([0-9]+(?:[\.,][0-9]{1,2})?)\s*(?:pesos|\$)?', line)
            if not m_price:
                continue
            
            # Buscar si tiene precio al final o al inicio
            # Probamos extraer el valor numérico
            price_str = m_price.group(1).replace(',', '.')
            try:
                precio = float(price_str)
            except ValueError:
                continue

            # El nombre es el resto de la línea
            nombre_part = line[:m_price.start()] + line[m_price.end():]
            nombre_part = re.sub(r'[:\-\$]', ' ', nombre_part).strip()

            # Extraer alias si los hay entre paréntesis ej: (alias: mila, milanga)
            aliases = []
            m_alias = re.search(r'\((?:alias|o|alias:)?\s*([^)]+)\)', nombre_part, re.IGNORECASE)
            if m_alias:
                aliases = [a.strip() for a in m_alias.group(1).split(",") if a.strip()]
                nombre_part = nombre_part[:m_alias.start()] + nombre_part[m_alias.end():]

            nombre = " ".join(nombre_part.split()).strip().capitalize()
            if not nombre:
                continue

            # Agregar nombre básico y variantes a alias
            norm_nom = normalize_text(nombre)
            if norm_nom not in aliases:
                aliases.append(norm_nom)

            new_items.append({
                "nombre": nombre,
                "precio": precio,
                "aliases": aliases
            })

        if new_items:
            self.set_items(new_items)
        return new_items

    def find_item(self, query: str):
        """Busca el plato más coincidente por nombre o alias."""
        q_norm = normalize_text(query)
        if not q_norm:
            return None

        # 1. Coincidencia exacta con algún alias o nombre
        for it in self.items:
            if normalize_text(it["nombre"]) == q_norm:
                return it
            for a in it.get("aliases", []):
                if normalize_text(a) == q_norm:
                    return it

        # 2. El alias o nombre está contenido en el query, o el query está contenido en el alias
        best_match = None
        best_len = 0
        for it in self.items:
            candidates = [it["nombre"]] + it.get("aliases", [])
            for c in candidates:
                c_norm = normalize_text(c)
                if c_norm in q_norm or q_norm in c_norm:
                    if len(c_norm) > best_len:
                        best_match = it
                        best_len = len(c_norm)

        # 3. Coincidencia por palabras individuales
        if not best_match:
            q_words = set(q_norm.split())
            max_common = 0
            for it in self.items:
                candidates = [it["nombre"]] + it.get("aliases", [])
                for c in candidates:
                    c_words = set(normalize_text(c).split())
                    common = len(q_words.intersection(c_words))
                    if common > max_common:
                        max_common = common
                        best_match = it

        return best_match

class GestorPedidos:
    def __init__(self, catalogo: Catalogo, filepath=PEDIDOS_FILE):
        self.catalogo = catalogo
        self.filepath = Path(filepath)
        self.pedidos = []
        self.load()

    def load(self):
        if self.filepath.exists():
            try:
                with open(self.filepath, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    self.pedidos = data.get("pedidos", [])
            except Exception as e:
                print(f"[!] Error al cargar pedidos: {e}")
                self.pedidos = []
        else:
            self.pedidos = []

    def save(self):
        with open(self.filepath, "w", encoding="utf-8") as f:
            json.dump({"pedidos": self.pedidos}, f, indent=2, ensure_ascii=False)

    def add_pedido(self, cliente: str, plato_query: str, dia: str, cantidad: int = 1, pagado: bool = False, notas: str = "", whatsapp: str = "", entregado: bool = False, usuario_id: int = None):
        item = self.catalogo.find_item(plato_query)
        if not item:
            # Si no está en el catálogo, guardamos el nombre tal como vino con precio 0 para advertir
            plato_nombre = plato_query
            precio_unitario = 0.0
            precio_encontrado = False
        else:
            if "(" in plato_query and ")" in plato_query:
                plato_nombre = plato_query
            else:
                plato_nombre = item["nombre"]
            precio_unitario = float(item["precio"])
            precio_encontrado = True

        next_id = max([p.get("id", 0) for p in self.pedidos], default=0) + 1
        dia_limpio = dia.strip().capitalize()
        # Normalizar tildes en días estándar
        norm_d = normalize_text(dia_limpio)
        if norm_d in DIAS_ORDEN:
            if norm_d == "miercoles":
                dia_limpio = "Miércoles"
            elif norm_d == "sabado":
                dia_limpio = "Sábado"

        total = round(cantidad * precio_unitario, 2)
        ahora = datetime.now()
        pedido = {
            "id": next_id,
            "usuario_id": int(usuario_id) if usuario_id else None,
            "dia": dia_limpio,
            "cliente": cliente.strip().title(),
            "whatsapp": whatsapp.strip(),
            "plato": plato_nombre,
            "cantidad": int(cantidad),
            "precio_unitario": precio_unitario,
            "total": total,
            "pagado": bool(pagado),
            "entregado": bool(entregado),
            "notas": notas.strip(),
            "fecha_registro": ahora.strftime("%Y-%m-%d %H:%M"),
            "created_at_iso": ahora.isoformat(),
            "precio_encontrado": precio_encontrado
        }
        self.pedidos.append(pedido)
        self.save()
        return pedido

    def toggle_pago(self, pedido_id: int):
        for p in self.pedidos:
            if p.get("id") == pedido_id:
                p["pagado"] = not p.get("pagado", False)
                self.save()
                return p
        return None

    def toggle_entrega(self, pedido_id: int):
        for p in self.pedidos:
            if p.get("id") == pedido_id:
                p["entregado"] = not p.get("entregado", False)
                self.save()
                return p
        return None

    def delete_pedido(self, pedido_id: int):
        target = None
        for p in self.pedidos:
            if p.get("id") == int(pedido_id):
                target = p
                break
        if target:
            self.pedidos = [p for p in self.pedidos if p.get("id") != int(pedido_id)]
            self.save()
            return target
        return None

    def get_pedidos_cliente(self, usuario_id=None, whatsapp=""):
        res = []
        wa_digits = re.sub(r'\D', '', str(whatsapp)) if whatsapp else ""
        for p in self.pedidos:
            match = False
            if usuario_id and p.get("usuario_id") == int(usuario_id):
                match = True
            elif wa_digits:
                p_digits = re.sub(r'\D', '', str(p.get("whatsapp", "")))
                if p_digits and p_digits == wa_digits:
                    match = True
            if match:
                res.append(p)
        return list(reversed(res))

    def parse_and_add_text(self, text: str, default_dia: str = ""):
        """
        Interpreta líneas o frases cotidianas.
        Formatos soportados:
        - 'Paula quiere pollo para el martes'
        - 'Paula pollo martes'
        - 'Paula 2 tartas de verdura jueves'
        - 'Lucas: 1 milanesa, pagado'
        - 'Martes: Juan pollo, Maria tarta'
        """
        lines = [l.strip() for l in text.strip().splitlines() if l.strip()]
        resultados = []

        dias_keywords = ["lunes", "martes", "miercoles", "miércoles", "jueves", "viernes", "sabado", "sábado", "domingo"]

        for raw_line in lines:
            # Si la línea empieza con un día: "Martes: Juan pollo, Maria tarta"
            match_header_dia = re.match(r'^(lunes|martes|miercoles|miércoles|jueves|viernes|sabado|sábado|domingo)\s*[:\-]\s*(.*)$', raw_line, re.IGNORECASE)
            dia_contexto = default_dia
            sub_items = [raw_line]
            if match_header_dia:
                dia_contexto = match_header_dia.group(1).capitalize()
                sub_items = [s.strip() for s in match_header_dia.group(2).split(",") if s.strip()]

            for item_str in sub_items:
                res = self._parse_single_order(item_str, dia_contexto, dias_keywords)
                if res:
                    ped = self.add_pedido(
                        cliente=res["cliente"],
                        plato_query=res["plato"],
                        dia=res["dia"],
                        cantidad=res["cantidad"],
                        pagado=res["pagado"],
                        notas=res.get("notas", "")
                    )
                    resultados.append(ped)

        return resultados

    def _parse_single_order(self, text: str, fallback_dia: str, dias_keywords: list):
        clean = text.strip()
        # Verificar si menciona "pagado" o "pago"
        pagado = False
        if re.search(r'\b(pagado|pago|ya pago|abono)\b', clean, re.IGNORECASE):
            pagado = True
            clean = re.sub(r'\b(pagado|pago|ya pago|abono)\b', '', clean, flags=re.IGNORECASE)

        # Buscar día en el texto
        dia = fallback_dia
        for d in dias_keywords:
            pattern = rf'\b(?:para\s+el\s+|el\s+)?({d})\b'
            m = re.search(pattern, clean, re.IGNORECASE)
            if m:
                dia = m.group(1).capitalize()
                # Remover la mención del día para facilitar extracción de plato
                clean = re.sub(pattern, '', clean, flags=re.IGNORECASE)
                break

        if not dia:
            dia = "Sin especificar"

        # Buscar cantidad: ej. "2 pollos", "3 tartas", "x2"
        cantidad = 1
        m_cant = re.search(r'\b(\d+)\s*(?:viandas?|porciones?|unidades?|de)?\b', clean, re.IGNORECASE)
        if m_cant:
            try:
                cantidad = int(m_cant.group(1))
                # quitar la cantidad del string
                clean = clean[:m_cant.start()] + clean[m_cant.end():]
            except ValueError:
                pass
        else:
            m_x = re.search(r'\bx\s*(\d+)\b', clean, re.IGNORECASE)
            if m_x:
                try:
                    cantidad = int(m_x.group(1))
                    clean = clean[:m_x.start()] + clean[m_x.end():]
                except ValueError:
                    pass

        # Quitar palabras de relleno: "quiere", "pide", "va a querer", "para", "de", "con" al inicio
        clean = re.sub(r'[:,-]', ' ', clean)
        words = clean.split()
        if not words:
            return None

        # Normalmente el cliente es la primera palabra o dos palabras antes de verbos como "quiere", "pide"
        m_verbo = re.search(r'\b(quiere|pide|pidio|va a querer|lleva|prefiere)\b', clean, re.IGNORECASE)
        if m_verbo:
            cliente = clean[:m_verbo.start()].strip()
            plato = clean[m_verbo.end():].strip()
        else:
            # Caso simple: "Paula pollo" o "Juan tarta de jamon"
            cliente = words[0]
            plato = " ".join(words[1:]).strip()

        # Limpiezas adicionales en el plato
        plato = re.sub(r'^(una|un|dos|tres|para)\s+', '', plato, flags=re.IGNORECASE).strip()

        return {
            "cliente": cliente if cliente else "Cliente",
            "plato": plato if plato else "Vianda",
            "dia": dia,
            "cantidad": max(1, cantidad),
            "pagado": pagado
        }

    def clear_all(self):
        self.pedidos = []
        self.save()

class GeneradorExcel:
    @staticmethod
    def export(catalogo: Catalogo, gestor: GestorPedidos, output_path=EXCEL_OUTPUT):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()

        # Estilos generales
        color_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Azul marino elegante
        color_sub_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")     # Azul medio
        color_accent_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")  # Azul pastel claro
        color_total_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")   # Amarillo suave
        color_green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")   # Verde claro
        color_red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")     # Rojo claro

        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_sub_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_bold = Font(name="Calibri", size=11, bold=True)
        font_regular = Font(name="Calibri", size=11)
        font_green = Font(name="Calibri", size=11, color="166534", bold=True)
        font_red = Font(name="Calibri", size=11, color="991B1B", bold=True)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        double_bottom_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='double', color='1E3A8A')
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # -------------------------------------------------------------
        # HOJA 1: PEDIDOS DETALLADOS
        # -------------------------------------------------------------
        ws_pedidos = wb.active
        ws_pedidos.title = "Pedidos"
        ws_pedidos.views.sheetView[0].showGridLines = True

        headers_pedidos = ["ID", "Día Entrega", "Cliente", "WhatsApp", "Plato Solicitado", "Cant.", "Precio Unit.", "Total ($)", "Estado Pago", "Entregado", "Notas / Fecha"]
        ws_pedidos.append(headers_pedidos)

        for col_idx in range(1, len(headers_pedidos) + 1):
            cell = ws_pedidos.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = color_header_fill
            cell.alignment = align_center
            cell.border = thin_border
        ws_pedidos.row_dimensions[1].height = 25

        row_num = 2
        for p in gestor.pedidos:
            ws_pedidos.append([
                p.get("id"),
                p.get("dia"),
                p.get("cliente"),
                p.get("whatsapp", ""),
                p.get("plato"),
                p.get("cantidad"),
                p.get("precio_unitario"),
                f"=F{row_num}*G{row_num}",
                "PAGADO" if p.get("pagado") else "PENDIENTE",
                "ENTREGADO" if p.get("entregado") else "PENDIENTE",
                p.get("notas") or p.get("fecha_registro", "")
            ])

            # Formatos de celda
            ws_pedidos.cell(row=row_num, column=1).alignment = align_center
            ws_pedidos.cell(row=row_num, column=2).alignment = align_center
            ws_pedidos.cell(row=row_num, column=3).alignment = align_left
            ws_pedidos.cell(row=row_num, column=4).alignment = align_center
            ws_pedidos.cell(row=row_num, column=5).alignment = align_left
            ws_pedidos.cell(row=row_num, column=6).alignment = align_center

            # Precio y total como moneda
            c_punit = ws_pedidos.cell(row=row_num, column=7)
            c_punit.number_format = '$#,##0.00'
            c_punit.alignment = align_right

            c_tot = ws_pedidos.cell(row=row_num, column=8)
            c_tot.number_format = '$#,##0.00'
            c_tot.alignment = align_right

            # Estado de pago
            c_pago = ws_pedidos.cell(row=row_num, column=9)
            c_pago.alignment = align_center
            if p.get("pagado"):
                c_pago.fill = color_green_fill
                c_pago.font = font_green
            else:
                c_pago.fill = color_red_fill
                c_pago.font = font_red

            # Estado de entrega
            c_ent = ws_pedidos.cell(row=row_num, column=10)
            c_ent.alignment = align_center
            if p.get("entregado"):
                c_ent.fill = color_green_fill
                c_ent.font = font_green
            else:
                c_ent.fill = color_accent_fill

            ws_pedidos.cell(row=row_num, column=11).alignment = align_left

            for c in range(1, len(headers_pedidos) + 1):
                ws_pedidos.cell(row=row_num, column=c).border = thin_border
            ws_pedidos.row_dimensions[row_num].height = 20
            row_num += 1

        # Fila de Totales
        if row_num > 2:
            total_row = row_num
            ws_pedidos.cell(row=total_row, column=5, value="TOTAL GENERAL:").font = font_bold
            ws_pedidos.cell(row=total_row, column=5).alignment = align_right
            
            c_sum_cant = ws_pedidos.cell(row=total_row, column=6, value=f"=SUM(F2:F{total_row-1})")
            c_sum_cant.font = font_bold
            c_sum_cant.alignment = align_center
            c_sum_cant.fill = color_total_fill

            c_sum_dinero = ws_pedidos.cell(row=total_row, column=8, value=f"=SUM(H2:H{total_row-1})")
            c_sum_dinero.font = font_bold
            c_sum_dinero.alignment = align_right
            c_sum_dinero.number_format = '$#,##0.00'
            c_sum_dinero.fill = color_total_fill

            for c in range(1, len(headers_pedidos) + 1):
                ws_pedidos.cell(row=total_row, column=c).border = double_bottom_border
            ws_pedidos.row_dimensions[total_row].height = 22

        # -------------------------------------------------------------
        # HOJA 2: PARA LA COCINA (RESUMEN POR DÍA)
        # -------------------------------------------------------------
        ws_cocina = wb.create_sheet(title="Para la Cocina")
        ws_cocina.views.sheetView[0].showGridLines = True

        headers_cocina = ["Día", "Plato a Preparar", "Cantidad a Cocinar", "Clientes que pidieron"]
        ws_cocina.append(headers_cocina)
        for col_idx in range(1, len(headers_cocina) + 1):
            cell = ws_cocina.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = color_sub_fill
            cell.alignment = align_center
            cell.border = thin_border
        ws_cocina.row_dimensions[1].height = 25

        # Agrupar pedidos por día y por plato
        dias_agrupados = {}
        for p in gestor.pedidos:
            d = p.get("dia", "Varios")
            plato = p.get("plato", "Sin nombre")
            cant = p.get("cantidad", 1)
            cli = p.get("cliente", "")
            dias_agrupados.setdefault(d, {}).setdefault(plato, {"cant": 0, "clientes": []})
            dias_agrupados[d][plato]["cant"] += cant
            dias_agrupados[d][plato]["clientes"].append(f"{cli} (x{cant})" if cant > 1 else cli)

        # Ordenar días
        def sort_dia_key(d):
            norm = normalize_text(d)
            return DIAS_ORDEN.get(norm, 99)

        row_c = 2
        for dia in sorted(dias_agrupados.keys(), key=sort_dia_key):
            platos_del_dia = dias_agrupados[dia]
            primera_fila = True
            for plato, info in platos_del_dia.items():
                detalle_clientes = ", ".join(info["clientes"])
                ws_cocina.append([
                    dia if primera_fila else "",
                    plato,
                    info["cant"],
                    detalle_clientes
                ])
                c_dia = ws_cocina.cell(row=row_c, column=1)
                c_dia.font = font_bold if primera_fila else font_regular
                c_dia.alignment = align_center
                if primera_fila:
                    c_dia.fill = color_accent_fill

                ws_cocina.cell(row=row_c, column=2).alignment = align_left
                c_cnt = ws_cocina.cell(row=row_c, column=3)
                c_cnt.alignment = align_center
                c_cnt.font = font_bold
                c_cnt.fill = color_total_fill
                ws_cocina.cell(row=row_c, column=4).alignment = align_left

                for col_idx in range(1, 5):
                    ws_cocina.cell(row=row_c, column=col_idx).border = thin_border
                ws_cocina.row_dimensions[row_c].height = 20
                row_c += 1
                primera_fila = False

        # -------------------------------------------------------------
        # HOJA 3: CUENTAS POR CLIENTE
        # -------------------------------------------------------------
        ws_clientes = wb.create_sheet(title="Cuentas Clientes")
        ws_clientes.views.sheetView[0].showGridLines = True

        headers_cli = ["Cliente", "WhatsApp", "Cant. Viandas", "Total a Pagar ($)", "Total Pagado ($)", "Saldo Pendiente ($)"]
        ws_clientes.append(headers_cli)
        for col_idx in range(1, len(headers_cli) + 1):
            cell = ws_clientes.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = color_header_fill
            cell.alignment = align_center
            cell.border = thin_border
        ws_clientes.row_dimensions[1].height = 25

        clientes_resumen = {}
        for p in gestor.pedidos:
            c = p.get("cliente", "Sin nombre")
            wa = p.get("whatsapp", "")
            cant = p.get("cantidad", 1)
            tot = p.get("total", 0.0)
            pag = tot if p.get("pagado") else 0.0

            clientes_resumen.setdefault(c, {"whatsapp": wa, "cant": 0, "total": 0.0, "pagado": 0.0})
            if wa and not clientes_resumen[c]["whatsapp"]:
                clientes_resumen[c]["whatsapp"] = wa
            clientes_resumen[c]["cant"] += cant
            clientes_resumen[c]["total"] += tot
            clientes_resumen[c]["pagado"] += pag

        row_cl = 2
        for cliente, data in sorted(clientes_resumen.items()):
            saldo = data["total"] - data["pagado"]
            ws_clientes.append([
                cliente,
                data.get("whatsapp", ""),
                data["cant"],
                data["total"],
                data["pagado"],
                f"=D{row_cl}-E{row_cl}"
            ])
            ws_clientes.cell(row=row_cl, column=1).alignment = align_left
            ws_clientes.cell(row=row_cl, column=2).alignment = align_center
            ws_clientes.cell(row=row_cl, column=3).alignment = align_center

            for c_idx in [4, 5, 6]:
                cell = ws_clientes.cell(row=row_cl, column=c_idx)
                cell.number_format = '$#,##0.00'
                cell.alignment = align_right

            c_saldo = ws_clientes.cell(row=row_cl, column=6)
            c_saldo.font = font_bold
            if saldo > 0:
                c_saldo.fill = color_red_fill
            else:
                c_saldo.fill = color_green_fill

            for col_idx in range(1, len(headers_cli) + 1):
                ws_clientes.cell(row=row_cl, column=col_idx).border = thin_border
            ws_clientes.row_dimensions[row_cl].height = 20
            row_cl += 1

        # -------------------------------------------------------------
        # HOJA 4: LISTA DE PRECIOS / MENÚ
        # -------------------------------------------------------------
        ws_menu = wb.create_sheet(title="Lista de Precios")
        ws_menu.views.sheetView[0].showGridLines = True

        headers_menu = ["ID", "Plato / Vianda", "Precio Vigente ($)", "Alias y Palabras Clave"]
        ws_menu.append(headers_menu)
        for col_idx in range(1, len(headers_menu) + 1):
            cell = ws_menu.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = color_sub_fill
            cell.alignment = align_center
            cell.border = thin_border
        ws_menu.row_dimensions[1].height = 25

        row_m = 2
        for item in catalogo.items:
            ws_menu.append([
                item.get("id"),
                item.get("nombre"),
                item.get("precio"),
                ", ".join(item.get("aliases", []))
            ])
            ws_menu.cell(row=row_m, column=1).alignment = align_center
            ws_menu.cell(row=row_m, column=2).alignment = align_left
            c_pr = ws_menu.cell(row=row_m, column=3)
            c_pr.number_format = '$#,##0.00'
            c_pr.alignment = align_right
            c_pr.font = font_bold
            ws_menu.cell(row=row_m, column=4).alignment = align_left

            for col_idx in range(1, len(headers_menu) + 1):
                ws_menu.cell(row=row_m, column=col_idx).border = thin_border
            ws_menu.row_dimensions[row_m].height = 20
            row_m += 1

        # Ajuste automático del ancho de columnas para todas las hojas
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val = cell.value
                    if val is not None:
                        s_val = str(val)
                        if s_val.startswith("="):
                            s_val = "$123,456.78"
                        max_len = max(max_len, len(s_val))
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Guardar en la ruta destino
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        print(f"[OK] Archivo Excel guardado en: {output_path}")

        # Copiar también al Escritorio si existe
        if DESKTOP_DIR.exists():
            try:
                shutil.copy2(output_path, DESKTOP_EXCEL)
                print(f"[OK] Copia actualizada en el Escritorio: {DESKTOP_EXCEL}")
            except Exception as e:
                print(f"[!] No se pudo copiar al Escritorio: {e}")

        return output_path

# =====================================================================
# CLI / Helper interactivo
# =====================================================================
def print_status(catalogo: Catalogo, gestor: GestorPedidos):
    print("=" * 60)
    print("📋 CATÁLOGO ACTUAL DE PRECIOS:")
    print("=" * 60)
    for it in catalogo.items:
        print(f"  • {it['nombre']}: ${it['precio']:,.2f}  (Alias: {', '.join(it.get('aliases', []))})")

    print("\n" + "=" * 60)
    print(f"📦 PEDIDOS REGISTRADOS ({len(gestor.pedidos)}):")
    print("=" * 60)
    if not gestor.pedidos:
        print("  (No hay pedidos cargados aún)")
    else:
        for p in gestor.pedidos:
            pago_str = "PAGADO" if p.get("pagado") else "PENDIENTE"
            print(f"  [{p['id']}] {p['dia']} | {p['cliente']} -> {p['cantidad']}x {p['plato']} = ${p['total']:,.2f} ({pago_str})")

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Motor de viandas y Excel")
    parser.add_argument("action", choices=["status", "parse", "excel", "clear", "add-dish", "set-menu"], nargs="?", default="status")
    parser.add_argument("--text", type=str, help="Texto o lista de pedidos o menú para procesar")
    parser.add_argument("--nombre", type=str, help="Nombre del plato para agregar/actualizar")
    parser.add_argument("--precio", type=float, help="Precio del plato")
    parser.add_argument("--aliases", type=str, help="Alias separados por comas")

    args = parser.parse_args()
    cat = Catalogo()
    gest = GestorPedidos(cat)

    if args.action == "status":
        print_status(cat, gest)
    elif args.action == "parse":
        if args.text:
            pedidos_agregados = gest.parse_and_add_text(args.text)
            print(f"[+] Se registraron {len(pedidos_agregados)} pedidos:")
            for p in pedidos_agregados:
                print(f"  - {p['cliente']} | {p['dia']}: {p['cantidad']}x {p['plato']} (${p['total']})")
            GeneradorExcel.export(cat, gest)
        else:
            print("Debes pasar el argumento --text")
    elif args.action == "set-menu":
        if args.text:
            items = cat.parse_and_set_menu_text(args.text)
            GeneradorExcel.export(cat, gest)
            print(f"[OK] Se cargaron {len(items)} platos al menú:")
            for it in items:
                print(f"  • {it['nombre']}: ${it['precio']} (alias: {', '.join(it['aliases'])})")
        else:
            print("Debes pasar el argumento --text con los platos y precios")
    elif args.action == "excel":
        GeneradorExcel.export(cat, gest)
    elif args.action == "clear":
        gest.clear_all()
        GeneradorExcel.export(cat, gest)
        print("[OK] Pedidos reseteados y Excel actualizado.")
    elif args.action == "add-dish":
        if args.nombre and args.precio is not None:
            aliases = [a.strip() for a in args.aliases.split(",")] if args.aliases else []
            cat.add_or_update_item(args.nombre, args.precio, aliases)
            GeneradorExcel.export(cat, gest)
            print(f"[OK] Plato '{args.nombre}' guardado con precio ${args.precio}.")

